"""
ragmine.ingest.parser
~~~~~~~~~~~~~~~~~~~~~
Extracts text from files. Docling if available, fallbacks otherwise.
"""

from __future__ import annotations

from pathlib import Path

from ragmine.core.config import Settings
from ragmine.core.protocols import ParsedDocument


TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".yaml", ".yml", ".toml", ".log", ".env"}
CODE_EXTENSIONS = {".py", ".java", ".js", ".ts", ".go", ".rs", ".rb", ".sh", ".sql", ".kt", ".xml"}
DOC_EXTENSIONS = {".pdf", ".docx", ".pptx", ".html", ".htm"}


class SmartParser:
    """Routes to the best parser per file type. Pluggable."""

    def __init__(self, settings: Settings | None = None):
        self._settings = settings

    def supports(self, file_path: str) -> bool:
        ext = Path(file_path).suffix.lower()
        return ext in TEXT_EXTENSIONS | CODE_EXTENSIONS | DOC_EXTENSIONS

    def parse(self, file_path: str) -> ParsedDocument:
        path = Path(file_path)
        ext = path.suffix.lower()

        if ext in TEXT_EXTENSIONS:
            return self._parse_text(path, source_type="text")

        if ext in CODE_EXTENSIONS:
            return self._parse_text(path, source_type="code")

        if ext in DOC_EXTENSIONS:
            return self._parse_document(path)

        # Best effort: try reading as text
        return self._parse_text(path, source_type="unknown")

    def _parse_text(self, path: Path, source_type: str) -> ParsedDocument:
        try:
            text = path.read_text(encoding="utf-8", errors="replace")
        except Exception as e:
            text = f"[Error reading file: {e}]"

        return ParsedDocument(
            text=text,
            source=str(path),
            source_type=source_type,
            metadata={"filename": path.name, "extension": path.suffix},
        )

    def _parse_document(self, path: Path) -> ParsedDocument:
        # Try Docling first (best quality)
        try:
            return self._parse_with_docling(path)
        except ImportError:
            pass

        # Fallback per type
        ext = path.suffix.lower()
        if ext == ".pdf":
            return self._parse_pdf_fallback(path)
        if ext == ".docx":
            return self._parse_docx_fallback(path)
        if ext in (".html", ".htm"):
            return self._parse_html_fallback(path)

        return self._parse_text(path, source_type="document")

    def _parse_with_docling(self, path: Path) -> ParsedDocument:
        from docling.document_converter import DocumentConverter

        converter = DocumentConverter()
        result = converter.convert(str(path))
        text = result.document.export_to_markdown()

        return ParsedDocument(
            text=text,
            source=str(path),
            source_type="document",
            metadata={"filename": path.name, "parser": "docling"},
        )

    def _parse_pdf_fallback(self, path: Path) -> ParsedDocument:
        try:
            import pdfplumber
            pages = []
            with pdfplumber.open(path) as pdf:
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        pages.append(t)
            text = "\n\n".join(pages)
        except ImportError:
            text = f"[Install pdfplumber or docling to parse PDFs: {path.name}]"

        return ParsedDocument(
            text=text, source=str(path), source_type="document",
            metadata={"filename": path.name, "parser": "pdfplumber"}, pages=pages if "pages" in dir() else None,
        )

    def _parse_docx_fallback(self, path: Path) -> ParsedDocument:
        try:
            from docx import Document
            doc = Document(str(path))
            text = "\n\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except ImportError:
            text = f"[Install python-docx or docling to parse DOCX: {path.name}]"

        return ParsedDocument(
            text=text, source=str(path), source_type="document",
            metadata={"filename": path.name, "parser": "python-docx"},
        )

    def _parse_html_fallback(self, path: Path) -> ParsedDocument:
        try:
            from html.parser import HTMLParser

            class TextExtractor(HTMLParser):
                def __init__(self):
                    super().__init__()
                    self.parts: list[str] = []

                def handle_data(self, data):
                    stripped = data.strip()
                    if stripped:
                        self.parts.append(stripped)

            raw = path.read_text(encoding="utf-8", errors="replace")
            extractor = TextExtractor()
            extractor.feed(raw)
            text = "\n".join(extractor.parts)
        except Exception:
            text = path.read_text(encoding="utf-8", errors="replace")

        return ParsedDocument(
            text=text, source=str(path), source_type="document",
            metadata={"filename": path.name, "parser": "html"},
        )
